from __future__ import annotations

import json
import time
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any

import streamlit as st
from PIL import Image, ImageChops, ImageEnhance
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# =========================================================
# 0. 프로젝트 경로
# =========================================================
# Nexis_LCD_Inspection/
# ├── data/
# │   ├── reference/
# │   └── capture/
# ├── results/
# │   ├── inspection_results.json
# │   └── manual_review_results.json   ← 자동 생성
# └── website/
#     └── web.py

WEBSITE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = WEBSITE_DIR.parent

REFERENCE_DIR = PROJECT_ROOT / "data" / "reference"
CAPTURE_DIR = PROJECT_ROOT / "data" / "capture"

INSPECTION_RESULT_PATH = (
    PROJECT_ROOT / "results" / "inspection_results.json"
)

MANUAL_REVIEW_PATH = (
    PROJECT_ROOT / "results" / "manual_review_results.json"
)

SUPPORTED_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".bmp",
    ".webp",
}


# =========================================================
# 1. Streamlit 설정
# =========================================================
st.set_page_config(
    page_title="NEXIS LCD Inspection",
    page_icon="🔍",
    layout="wide",
    initial_sidebar_state="collapsed",
)


# =========================================================
# 2. CSS
# =========================================================
st.markdown(
    """
    <style>
    .block-container {
        max-width: 1500px;
        padding-top: 1.4rem;
        padding-bottom: 3rem;
    }

    .hero {
        padding: 2.3rem 2.5rem;
        border: 1px solid #e5e7eb;
        border-radius: 22px;
        background:
            linear-gradient(135deg, #ffffff 0%, #f4f7fb 100%);
        margin-bottom: 1.3rem;
    }

    .hero-title {
        font-size: 2.35rem;
        font-weight: 850;
        color: #111827;
        margin-bottom: 0.35rem;
    }

    .hero-subtitle {
        font-size: 1rem;
        color: #6b7280;
    }

    .section-title {
        font-size: 1.35rem;
        font-weight: 800;
        color: #111827;
        margin-top: 1rem;
        margin-bottom: 0.8rem;
    }

    .status-card {
        border-radius: 18px;
        padding: 1.4rem 1.6rem;
        text-align: center;
        font-size: 2rem;
        font-weight: 850;
        border: 1px solid rgba(0, 0, 0, 0.08);
        margin-bottom: 1rem;
    }

    .status-pass {
        background: rgba(34, 197, 94, 0.12);
        color: #15803d;
    }

    .status-fail {
        background: rgba(239, 68, 68, 0.12);
        color: #b91c1c;
    }

    .status-review {
        background: rgba(245, 158, 11, 0.16);
        color: #b45309;
    }

    .status-unknown {
        background: rgba(107, 114, 128, 0.12);
        color: #4b5563;
    }

    .info-box {
        border: 1px solid #e5e7eb;
        background: #f8fafc;
        border-radius: 14px;
        padding: 1rem 1.15rem;
        margin-bottom: 0.8rem;
    }

    .reason-card {
        border: 1px solid #e5e7eb;
        border-left: 5px solid #9ca3af;
        border-radius: 12px;
        background: #ffffff;
        padding: 0.9rem 1rem;
        margin-bottom: 0.7rem;
    }

    .caption {
        color: #6b7280;
        font-size: 0.86rem;
        margin-top: 0.3rem;
        word-break: break-all;
    }

    div[data-testid="stMetric"] {
        border: 1px solid #e5e7eb;
        border-radius: 15px;
        padding: 1rem;
        background: #ffffff;
    }

    div[data-testid="stMetricLabel"] {
        font-weight: 700;
    }

    .stButton > button,
    .stDownloadButton > button {
        border-radius: 12px;
        font-weight: 750;
        min-height: 46px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 3. 공통 함수
# =========================================================
def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"파일이 없습니다: {path}")

    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)

    if not isinstance(data, dict):
        raise ValueError(f"JSON 최상위 구조가 객체가 아닙니다: {path}")

    return data


def save_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    temporary_path = path.with_suffix(path.suffix + ".tmp")

    with temporary_path.open("w", encoding="utf-8") as file:
        json.dump(
            data,
            file,
            ensure_ascii=False,
            indent=2,
        )

    temporary_path.replace(path)


def load_manual_reviews() -> dict[str, Any]:
    if not MANUAL_REVIEW_PATH.exists():
        return {"reviews": {}}

    try:
        data = load_json(MANUAL_REVIEW_PATH)
    except (json.JSONDecodeError, ValueError):
        return {"reviews": {}}

    reviews = data.get("reviews")

    if not isinstance(reviews, dict):
        data["reviews"] = {}

    return data


def save_manual_review(
    file_name: str,
    decision: str,
    original_status: str,
) -> None:
    data = load_manual_reviews()
    reviews = data.setdefault("reviews", {})

    reviews[file_name] = {
        "manual_decision": decision,
        "original_status": original_status,
        "reviewed_at": time.strftime("%Y-%m-%d %H:%M:%S"),
    }

    save_json(MANUAL_REVIEW_PATH, data)


def find_image_files(folder: Path) -> list[Path]:
    if not folder.exists():
        return []

    return sorted(
        [
            path
            for path in folder.iterdir()
            if path.is_file()
            and path.suffix.lower() in SUPPORTED_EXTENSIONS
        ],
        key=lambda path: path.name.lower(),
    )


def build_image_map(files: list[Path]) -> dict[str, Path]:
    return {
        path.stem: path
        for path in files
    }


def normalize_status(value: Any) -> str:
    status = str(value or "UNKNOWN").strip().upper()

    if status in {"PASS", "FAIL", "REVIEW"}:
        return status

    return "UNKNOWN"


def get_results_dict(
    inspection_data: dict[str, Any],
) -> dict[str, dict[str, Any]]:
    results = inspection_data.get("results", {})

    if not isinstance(results, dict):
        raise ValueError(
            "inspection_results.json의 'results'는 객체여야 합니다."
        )

    normalized_results: dict[str, dict[str, Any]] = {}

    for file_name, result in results.items():
        if isinstance(result, dict):
            normalized_results[str(file_name)] = result

    return normalized_results


def get_status_file_map(
    results: dict[str, dict[str, Any]],
) -> dict[str, list[str]]:
    status_map = {
        "PASS": [],
        "FAIL": [],
        "REVIEW": [],
        "UNKNOWN": [],
    }

    for file_name, result in results.items():
        status = normalize_status(result.get("final_status"))
        status_map.setdefault(status, []).append(file_name)

    for file_names in status_map.values():
        file_names.sort(key=str.lower)

    return status_map


def open_image(path: Path) -> Image.Image:
    with Image.open(path) as image:
        return image.convert("RGB")


def create_difference_image(
    reference_path: Path,
    capture_path: Path,
    enhancement_factor: float = 3.0,
) -> Image.Image:
    reference = open_image(reference_path)
    capture = open_image(capture_path)

    if capture.size != reference.size:
        capture = capture.resize(reference.size)

    difference = ImageChops.difference(
        reference,
        capture,
    )

    return ImageEnhance.Contrast(
        difference
    ).enhance(
        enhancement_factor
    )


def resolve_image_path(
    file_name: str,
    image_map: dict[str, Path],
) -> Path | None:
    return image_map.get(Path(file_name).stem)


def get_status_class(status: str) -> str:
    return {
        "PASS": "status-pass",
        "FAIL": "status-fail",
        "REVIEW": "status-review",
    }.get(status, "status-unknown")


def display_status_card(
    status: str,
    file_name: str,
) -> None:
    icon = {
        "PASS": "✅",
        "FAIL": "❌",
        "REVIEW": "⚠️",
    }.get(status, "❔")

    st.markdown(
        f"""
        <div class="status-card {get_status_class(status)}">
            {icon} {escape(status)}
            <div style="
                font-size: 0.92rem;
                font-weight: 550;
                margin-top: 0.35rem;
            ">
                {escape(file_name)}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def get_final_reasons(
    result: dict[str, Any],
) -> list[str]:
    raw_reasons = result.get("final_reasons", [])

    if isinstance(raw_reasons, list):
        return [
            str(reason)
            for reason in raw_reasons
            if str(reason).strip()
        ]

    if raw_reasons:
        return [str(raw_reasons)]

    return []


def calculate_system_reliability(
    results: dict[str, dict[str, Any]],
) -> tuple[float | None, int, int]:
    """
    현재 신뢰도 정의:
    expected_result와 final_status가 모두 있는 항목 중
    두 판정이 정확히 일치한 비율.
    """
    comparable_count = 0
    matched_count = 0

    for result in results.values():
        expected = normalize_status(
            result.get("expected_result")
        )
        actual = normalize_status(
            result.get("final_status")
        )

        if expected == "UNKNOWN" or actual == "UNKNOWN":
            continue

        comparable_count += 1

        if expected == actual:
            matched_count += 1

    if comparable_count == 0:
        return None, matched_count, comparable_count

    reliability = (
        matched_count
        / comparable_count
        * 100
    )

    return reliability, matched_count, comparable_count


def status_korean(status: str) -> str:
    return {
        "PASS": "합격",
        "FAIL": "불합격",
        "REVIEW": "검토 필요",
        "UNKNOWN": "판정 없음",
    }.get(status, status)


def add_file_list_section(
    document: Document,
    title: str,
    file_names: list[str],
) -> None:
    """
    상태별 파일명 목록을 워드 표로 추가한다.
    """
    document.add_heading(
        f"{title} ({len(file_names)}개)",
        level=2,
    )

    if not file_names:
        document.add_paragraph(
            "해당 상태의 파일이 없습니다."
        )
        return

    table = document.add_table(
        rows=1,
        cols=2,
    )
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    header_cells[0].text = "번호"
    header_cells[1].text = "파일명"

    for index, file_name in enumerate(
        file_names,
        start=1,
    ):
        row_cells = table.add_row().cells
        row_cells[0].text = str(index)
        row_cells[1].text = file_name


def build_word_report(
    inspection_data: dict[str, Any],
    inspection_results: dict[str, dict[str, Any]],
    elapsed_time: float | None,
    reliability: float | None,
    matched_count: int,
    comparable_count: int,
) -> bytes:
    """
    사람이 바로 읽을 수 있는 최종 검사 리포트를
    Word(.docx) 형식으로 생성한다.
    """
    status_map = get_status_file_map(
        inspection_results
    )

    manual_reviews_data = load_manual_reviews()
    manual_reviews = manual_reviews_data.get(
        "reviews",
        {},
    )

    if not isinstance(manual_reviews, dict):
        manual_reviews = {}

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        "NEXIS LCD 자동 판독 최종 결과 리포트"
    )
    title_run.bold = True
    title_run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Reference·Capture 이미지 자동 검사 결과 및 REVIEW 수동 판정 기록"
    )

    document.add_paragraph()

    # 1. 검사 개요
    document.add_heading(
        "1. 검사 개요",
        level=1,
    )

    overview_table = document.add_table(
        rows=0,
        cols=2,
    )
    overview_table.style = "Table Grid"

    overview_items = [
        (
            "리포트 생성 시각",
            time.strftime("%Y-%m-%d %H:%M:%S"),
        ),
        (
            "총 검사 파일 수",
            str(len(inspection_results)),
        ),
        (
            "총 소요 시간",
            (
                f"{elapsed_time:.2f}초"
                if isinstance(
                    elapsed_time,
                    (int, float),
                )
                else "확인 불가"
            ),
        ),
        (
            "시스템 신뢰도",
            (
                f"{reliability:.2f}%"
                if isinstance(
                    reliability,
                    (int, float),
                )
                else "계산 불가"
            ),
        ),
        (
            "신뢰도 검증 데이터",
            f"{matched_count}개 일치 / {comparable_count}개 비교",
        ),
    ]

    for label, value in overview_items:
        row_cells = overview_table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = value

    # 2. 전체 판정 요약
    document.add_heading(
        "2. 전체 판정 요약",
        level=1,
    )

    summary_table = document.add_table(
        rows=1,
        cols=4,
    )
    summary_table.style = "Table Grid"

    summary_headers = [
        "총 검사",
        "PASS",
        "REVIEW",
        "FAIL",
    ]

    summary_values = [
        len(inspection_results),
        len(status_map["PASS"]),
        len(status_map["REVIEW"]),
        len(status_map["FAIL"]),
    ]

    for index, header in enumerate(
        summary_headers
    ):
        summary_table.rows[0].cells[index].text = (
            header
        )

    summary_row = summary_table.add_row().cells

    for index, value in enumerate(
        summary_values
    ):
        summary_row[index].text = str(value)

    document.add_paragraph(
        "PASS는 자동 판독 기준을 충족한 항목, "
        "FAIL은 자동 판독 기준에서 불합격한 항목, "
        "REVIEW는 사람이 추가로 확인해야 하는 항목입니다."
    )

    # 3. 상태별 파일 목록
    document.add_heading(
        "3. 상태별 파일 목록",
        level=1,
    )

    add_file_list_section(
        document,
        "PASS 파일",
        status_map["PASS"],
    )

    add_file_list_section(
        document,
        "REVIEW 파일",
        status_map["REVIEW"],
    )

    add_file_list_section(
        document,
        "FAIL 파일",
        status_map["FAIL"],
    )

    # 4. REVIEW 수동 확정 결과
    document.add_heading(
        "4. REVIEW 수동 확정 결과",
        level=1,
    )

    if manual_reviews:
        manual_table = document.add_table(
            rows=1,
            cols=4,
        )
        manual_table.style = "Table Grid"

        headers = [
            "파일명",
            "자동 판정",
            "사람의 최종 판정",
            "판정 시각",
        ]

        for index, header in enumerate(
            headers
        ):
            manual_table.rows[0].cells[index].text = (
                header
            )

        for file_name in sorted(
            manual_reviews.keys(),
            key=str.lower,
        ):
            review = manual_reviews[file_name]

            if not isinstance(review, dict):
                continue

            row_cells = manual_table.add_row().cells
            row_cells[0].text = file_name
            row_cells[1].text = status_korean(
                normalize_status(
                    review.get("original_status")
                )
            )
            row_cells[2].text = status_korean(
                normalize_status(
                    review.get("manual_decision")
                )
            )
            row_cells[3].text = str(
                review.get(
                    "reviewed_at",
                    "-",
                )
            )
    else:
        document.add_paragraph(
            "아직 사람이 확정한 REVIEW 항목이 없습니다."
        )

    # 5. 파일별 세부 결과
    document.add_heading(
        "5. 파일별 세부 판독 결과",
        level=1,
    )

    for file_name in sorted(
        inspection_results.keys(),
        key=str.lower,
    ):
        result = inspection_results[file_name]
        status = normalize_status(
            result.get("final_status")
        )

        document.add_heading(
            file_name,
            level=2,
        )

        detail_table = document.add_table(
            rows=0,
            cols=2,
        )
        detail_table.style = "Table Grid"

        detail_items = [
            (
                "자동 최종 판정",
                status_korean(status),
            ),
            (
                "화면 ID",
                str(
                    result.get(
                        "screen_id",
                        Path(file_name).stem,
                    )
                ),
            ),
            (
                "카테고리",
                str(
                    result.get(
                        "category",
                        "-",
                    )
                ),
            ),
            (
                "프로파일",
                str(
                    result.get(
                        "profile",
                        "-",
                    )
                ),
            ),
            (
                "예상 결과",
                status_korean(
                    normalize_status(
                        result.get(
                            "expected_result"
                        )
                    )
                ),
            ),
        ]

        manual_result = manual_reviews.get(
            file_name
        )

        if isinstance(manual_result, dict):
            detail_items.append(
                (
                    "사람의 최종 판정",
                    status_korean(
                        normalize_status(
                            manual_result.get(
                                "manual_decision"
                            )
                        )
                    ),
                )
            )

        for label, value in detail_items:
            row_cells = detail_table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = value

        reasons = get_final_reasons(result)

        document.add_paragraph(
            "판정 사유:",
            style=None,
        )

        if reasons:
            for reason in reasons:
                document.add_paragraph(
                    reason,
                    style="List Bullet",
                )
        else:
            document.add_paragraph(
                "기록된 세부 오류 사유가 없습니다.",
                style="List Bullet",
            )

        roi_summary = result.get(
            "roi_decision_summary",
            {},
        )

        if isinstance(roi_summary, dict):
            document.add_paragraph(
                "ROI 판정 요약:"
            )

            roi_table = document.add_table(
                rows=1,
                cols=4,
            )
            roi_table.style = "Table Grid"

            roi_headers = [
                "PASS",
                "REVIEW",
                "FAIL",
                "누락 문자 확정",
            ]

            roi_values = [
                roi_summary.get(
                    "pass_count",
                    "-",
                ),
                roi_summary.get(
                    "review_count",
                    "-",
                ),
                roi_summary.get(
                    "fail_count",
                    "-",
                ),
                roi_summary.get(
                    "confirmed_missing_text_count",
                    "-",
                ),
            ]

            for index, header in enumerate(
                roi_headers
            ):
                roi_table.rows[0].cells[index].text = (
                    header
                )

            roi_row = roi_table.add_row().cells

            for index, value in enumerate(
                roi_values
            ):
                roi_row[index].text = str(value)

        document.add_paragraph()

    # 6. 신뢰도 설명
    document.add_heading(
        "6. 시스템 신뢰도 산정 기준",
        level=1,
    )

    document.add_paragraph(
        "현재 시스템 신뢰도는 각 이미지의 expected_result와 "
        "자동 판독 결과 final_status가 모두 존재하는 항목을 대상으로, "
        "두 값이 정확히 일치한 비율로 계산합니다."
    )

    if reliability is not None:
        document.add_paragraph(
            f"계산 결과: {matched_count}개 일치 / "
            f"{comparable_count}개 비교 = "
            f"{reliability:.2f}%"
        )
    else:
        document.add_paragraph(
            "비교 가능한 expected_result 데이터가 없어 "
            "신뢰도를 계산하지 못했습니다."
        )

    # 원본 summary 보존
    original_summary = inspection_data.get(
        "summary",
        {},
    )

    if isinstance(original_summary, dict):
        document.add_heading(
            "7. 원본 분석 요약 데이터",
            level=1,
        )

        for key, value in original_summary.items():
            document.add_paragraph(
                f"{key}: {value}"
            )

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def open_detail_page(file_name: str) -> None:
    """
    상태별 파일 목록에서 파일명을 누르면
    세부 판독 화면으로 이동한다.
    """
    st.session_state.selected_detail_file = file_name
    st.session_state.current_view = "detail"
    st.rerun()


def render_file_buttons(
    title: str,
    file_names: list[str],
    status: str,
) -> None:
    """
    상태별 파일명 버튼을 표시한다.

    REVIEW 파일 중 사람이 PASS/FAIL 수동 판정을 완료한 경우:
        파일명 (수동 처리 완료: PASS)
        파일명 (수동 처리 완료: FAIL)
    형식으로 표시한다.
    """
    st.markdown(f"#### {title}")

    if not file_names:
        st.info("해당 파일이 없습니다.")
        return

    manual_reviews_data = load_manual_reviews()
    manual_reviews = manual_reviews_data.get(
        "reviews",
        {},
    )

    if not isinstance(manual_reviews, dict):
        manual_reviews = {}

    for index, file_name in enumerate(
        file_names,
        start=1,
    ):
        button_label = file_name

        manual_review = manual_reviews.get(
            file_name
        )

        if (
            status == "REVIEW"
            and isinstance(
                manual_review,
                dict,
            )
        ):
            manual_decision = normalize_status(
                manual_review.get(
                    "manual_decision"
                )
            )

            if manual_decision in {
                "PASS",
                "FAIL",
            }:
                button_label = (
                    f"{file_name} "
                    f"(수동 처리 완료: {manual_decision})"
                )

        if st.button(
            button_label,
            key=f"open_{status}_{index}_{file_name}",
            use_container_width=True,
        ):
            open_detail_page(file_name)


def render_roi_details(roi_decisions: Any) -> None:
    if not isinstance(
        roi_decisions,
        list,
    ) or not roi_decisions:
        st.info(
            "ROI별 세부 판정 데이터가 없습니다."
        )
        return

    for index, roi in enumerate(
        roi_decisions,
        start=1,
    ):
        if not isinstance(roi, dict):
            continue

        roi_id = str(
            roi.get(
                "roi_id",
                roi.get("id", f"ROI-{index}"),
            )
        )

        roi_status = normalize_status(
            roi.get(
                "roi_final_status",
                roi.get(
                    "final_status",
                    roi.get("status"),
                ),
            )
        )

        with st.expander(
            f"{index}. {roi_id} — {roi_status}",
            expanded=(roi_status != "PASS"),
        ):
            st.json(roi)


# =========================================================
# 4. 초기 상태
# =========================================================
if "inspection_started" not in st.session_state:
    st.session_state.inspection_started = False

if "inspection_finished" not in st.session_state:
    st.session_state.inspection_finished = False

if "elapsed_time_sec" not in st.session_state:
    st.session_state.elapsed_time_sec = None

if "selected_detail_file" not in st.session_state:
    st.session_state.selected_detail_file = None

if "current_view" not in st.session_state:
    st.session_state.current_view = "summary"


# =========================================================
# 5. 데이터 불러오기
# =========================================================
try:
    inspection_data = load_json(
        INSPECTION_RESULT_PATH
    )

    inspection_results = get_results_dict(
        inspection_data
    )

except (
    FileNotFoundError,
    json.JSONDecodeError,
    ValueError,
) as error:
    st.error(str(error))
    st.stop()


reference_files = find_image_files(
    REFERENCE_DIR
)

capture_files = find_image_files(
    CAPTURE_DIR
)

reference_map = build_image_map(
    reference_files
)

capture_map = build_image_map(
    capture_files
)


# =========================================================
# 6. 헤더
# =========================================================
st.markdown(
    """
    <div class="hero">
        <div class="hero-title">
            NEXIS LCD Inspection System
        </div>
        <div class="hero-subtitle">
            LCD Reference·Capture 이미지 자동 판독 및
            REVIEW 수동 확정 시스템
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# 7. 초기 화면
# =========================================================
if not st.session_state.inspection_started:
    st.markdown(
        '<div class="section-title">검사 준비</div>',
        unsafe_allow_html=True,
    )

    info_col1, info_col2, info_col3 = st.columns(3)

    info_col1.metric(
        "Reference 이미지",
        len(reference_files),
    )

    info_col2.metric(
        "Capture 이미지",
        len(capture_files),
    )

    info_col3.metric(
        "판독 결과 데이터",
        len(inspection_results),
    )

    st.markdown(
        """
        <div class="info-box">
            판독 시작 버튼을 누르면 현재 연결된
            <b>inspection_results.json</b>을 읽고,
            전체 결과를 순차적으로 불러옵니다.
        </div>
        """,
        unsafe_allow_html=True,
    )

    if st.button(
        "▶ 판독 시작",
        type="primary",
        use_container_width=True,
    ):
        st.session_state.inspection_started = True
        st.session_state.inspection_finished = False
        st.session_state.current_view = "summary"

        start_time = time.perf_counter()

        progress_bar = st.progress(0)
        progress_text = st.empty()
        current_file_text = st.empty()

        file_names = list(
            inspection_results.keys()
        )

        total_count = len(file_names)

        for index, file_name in enumerate(
            file_names,
            start=1,
        ):
            progress = (
                index / total_count
                if total_count
                else 1.0
            )

            progress_bar.progress(progress)

            progress_text.markdown(
                f"**판독 진행률:** "
                f"{index} / {total_count} "
                f"({progress * 100:.1f}%)"
            )

            current_file_text.caption(
                f"현재 처리 중: {file_name}"
            )

            time.sleep(0.025)

        st.session_state.elapsed_time_sec = (
            time.perf_counter()
            - start_time
        )

        st.session_state.inspection_finished = True
        st.rerun()

    st.stop()


# =========================================================
# 8. 결과 공통 계산
# =========================================================
status_file_map = get_status_file_map(
    inspection_results
)

pass_files = status_file_map["PASS"]
fail_files = status_file_map["FAIL"]
review_files = status_file_map["REVIEW"]

elapsed_time = st.session_state.elapsed_time_sec

reliability, matched_count, comparable_count = (
    calculate_system_reliability(
        inspection_results
    )
)


# =========================================================
# 9. 세부 판독 화면
# =========================================================
if st.session_state.current_view == "detail":
    selected_file_name = (
        st.session_state.selected_detail_file
    )

    if selected_file_name not in inspection_results:
        st.error(
            "선택한 파일의 결과를 찾지 못했습니다."
        )

        if st.button(
            "← 결과 요약 화면으로 돌아가기"
        ):
            st.session_state.current_view = "summary"
            st.rerun()

        st.stop()

    top_back_col, top_title_col = st.columns(
        [1, 4]
    )

    with top_back_col:
        if st.button(
            "← 결과 목록",
            use_container_width=True,
        ):
            st.session_state.current_view = "summary"
            st.rerun()

    with top_title_col:
        st.markdown(
            '<div class="section-title">'
            '세부 이미지 판독'
            '</div>',
            unsafe_allow_html=True,
        )

    selected_result = inspection_results[
        selected_file_name
    ]

    selected_status = normalize_status(
        selected_result.get("final_status")
    )

    selected_reference_path = resolve_image_path(
        selected_file_name,
        reference_map,
    )

    selected_capture_path = resolve_image_path(
        selected_file_name,
        capture_map,
    )

    display_status_card(
        selected_status,
        selected_file_name,
    )

    # -----------------------------------------------------
    # 이미지 표시
    # -----------------------------------------------------
    reference_col, capture_col = st.columns(2)

    with reference_col:
        st.markdown("### Reference")

        if selected_reference_path:
            st.image(
                str(selected_reference_path),
                use_container_width=True,
            )

            st.markdown(
                f"""
                <div class="caption">
                    {escape(selected_reference_path.name)}
                </div>
                """,
                unsafe_allow_html=True,
            )
        else:
            st.error(
                "Reference 이미지를 찾지 못했습니다."
            )

    with capture_col:
        st.markdown("### Capture")

        if selected_capture_path:
            st.image(
                str(selected_capture_path),
                use_container_width=True,
            )

            st.markdown(
                f"""
                <div class="caption">
                    {escape(selected_capture_path.name)}
                </div>
                """,
                unsafe_allow_html=True,
            )
        else:
            st.error(
                "Capture 이미지를 찾지 못했습니다."
            )

    # -----------------------------------------------------
    # REVIEW 수동 판정 버튼: 사진 바로 아래
    # -----------------------------------------------------
    manual_review_data = load_manual_reviews()
    manual_reviews = manual_review_data.get(
        "reviews",
        {},
    )

    existing_manual_review = (
        manual_reviews.get(selected_file_name)
        if isinstance(manual_reviews, dict)
        else None
    )

    if selected_status == "REVIEW":
        st.markdown(
            '<div class="section-title">'
            '사람의 최종 판정'
            '</div>',
            unsafe_allow_html=True,
        )

        st.warning(
            "Reference와 Capture 이미지를 비교한 뒤 "
            "PASS 또는 FAIL로 확정하세요."
        )

        manual_pass_col, manual_fail_col = st.columns(2)

        with manual_pass_col:
            if st.button(
                "✅ PASS로 확정",
                type="primary",
                use_container_width=True,
            ):
                save_manual_review(
                    file_name=selected_file_name,
                    decision="PASS",
                    original_status=selected_status,
                )

                st.success(
                    f"{selected_file_name}을 PASS로 저장했습니다."
                )
                st.rerun()

        with manual_fail_col:
            if st.button(
                "❌ FAIL로 확정",
                use_container_width=True,
            ):
                save_manual_review(
                    file_name=selected_file_name,
                    decision="FAIL",
                    original_status=selected_status,
                )

                st.success(
                    f"{selected_file_name}을 FAIL로 저장했습니다."
                )
                st.rerun()

        if isinstance(existing_manual_review, dict):
            st.info(
                "현재 저장된 수동 판정: "
                f"**{existing_manual_review.get('manual_decision', '-')}** "
                f"({existing_manual_review.get('reviewed_at', '-')})"
            )

    # -----------------------------------------------------
    # Difference 이미지
    # -----------------------------------------------------
    st.markdown(
        '<div class="section-title">'
        'Difference'
        '</div>',
        unsafe_allow_html=True,
    )

    if (
        selected_reference_path
        and selected_capture_path
    ):
        try:
            difference_image = create_difference_image(
                selected_reference_path,
                selected_capture_path,
            )

            st.image(
                difference_image,
                use_container_width=True,
            )
        except Exception as error:
            st.error(
                "Difference 이미지를 생성하지 못했습니다."
            )
            st.code(str(error))
    else:
        st.info(
            "두 이미지가 모두 있어야 차이 이미지를 표시할 수 있습니다."
        )

    # -----------------------------------------------------
    # 세부 오류 사항
    # -----------------------------------------------------
    st.markdown(
        '<div class="section-title">'
        '세부 오류 사항'
        '</div>',
        unsafe_allow_html=True,
    )

    final_reasons = get_final_reasons(
        selected_result
    )

    if final_reasons:
        for index, reason in enumerate(
            final_reasons,
            start=1,
        ):
            st.markdown(
                f"""
                <div class="reason-card">
                    <b>{index}.</b>
                    {escape(reason)}
                </div>
                """,
                unsafe_allow_html=True,
            )
    else:
        st.success(
            "기록된 최종 오류 사유가 없습니다."
        )

    detail_col1, detail_col2 = st.columns(2)

    diff_summary = selected_result.get(
        "diff_summary",
        {},
    )

    roi_summary = selected_result.get(
        "roi_decision_summary",
        {},
    )

    with detail_col1:
        st.markdown("#### 차이 분석 요약")

        if isinstance(diff_summary, dict):
            st.json(diff_summary)
        else:
            st.info(
                "차이 분석 요약 데이터가 없습니다."
            )

    with detail_col2:
        st.markdown("#### ROI 판정 집계")

        if isinstance(roi_summary, dict):
            st.json(roi_summary)
        else:
            st.info(
                "ROI 판정 집계 데이터가 없습니다."
            )

    st.markdown("#### ROI별 상세 판정")

    render_roi_details(
        selected_result.get(
            "roi_decisions",
            [],
        )
    )

    with st.expander(
        "현재 선택 결과 원본 JSON",
        expanded=False,
    ):
        st.json(selected_result)

    st.stop()


# =========================================================
# 10. 결과 요약 화면
# =========================================================
completion_col, download_col = st.columns(
    [4, 1.4]
)

with completion_col:
    st.success(
        "전체 판독 결과 불러오기가 완료되었습니다."
    )

word_report_bytes = build_word_report(
    inspection_data=inspection_data,
    inspection_results=inspection_results,
    elapsed_time=elapsed_time,
    reliability=reliability,
    matched_count=matched_count,
    comparable_count=comparable_count,
)

with download_col:
    st.download_button(
        label="⬇ 최종 결과 리포트",
        data=word_report_bytes,
        file_name=(
            "NEXIS_LCD_Inspection_Report_"
            + time.strftime("%Y%m%d_%H%M%S")
            + ".docx"
        ),
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        use_container_width=True,
    )


summary_col1, summary_col2, summary_col3, summary_col4 = (
    st.columns(4)
)

summary_col1.metric(
    "총 소요 시간",
    (
        f"{elapsed_time:.2f} sec"
        if isinstance(elapsed_time, (int, float))
        else "-"
    ),
)

summary_col2.metric(
    "PASS",
    len(pass_files),
)

summary_col3.metric(
    "REVIEW",
    len(review_files),
)

summary_col4.metric(
    "FAIL",
    len(fail_files),
)


# =========================================================
# 11. 판정별 파일 목록
# =========================================================
st.markdown(
    '<div class="section-title">'
    '상태별 사진 파일명 목록'
    '</div>',
    unsafe_allow_html=True,
)

pass_tab, review_tab, fail_tab = st.tabs(
    [
        f"PASS ({len(pass_files)})",
        f"REVIEW ({len(review_files)})",
        f"FAIL ({len(fail_files)})",
    ]
)

with pass_tab:
    render_file_buttons(
        "PASS 파일",
        pass_files,
        "PASS",
    )

with review_tab:
    render_file_buttons(
        "REVIEW 파일",
        review_files,
        "REVIEW",
    )

with fail_tab:
    render_file_buttons(
        "FAIL 파일",
        fail_files,
        "FAIL",
    )


# =========================================================
# 12. 시스템 신뢰도
# =========================================================
st.divider()

st.markdown(
    '<div class="section-title">'
    '전체 시스템 신뢰도'
    '</div>',
    unsafe_allow_html=True,
)

if reliability is None:
    st.warning(
        "expected_result가 없어 시스템 신뢰도를 계산할 수 없습니다."
    )
else:
    reliability_col1, reliability_col2, reliability_col3 = (
        st.columns(3)
    )

    reliability_col1.metric(
        "시스템 신뢰도",
        f"{reliability:.2f}%",
    )

    reliability_col2.metric(
        "예상 결과와 일치",
        matched_count,
    )

    reliability_col3.metric(
        "검증 가능 데이터",
        comparable_count,
    )

    st.progress(
        min(
            max(
                reliability / 100,
                0.0,
            ),
            1.0,
        )
    )

    st.caption(
        "현재 신뢰도는 expected_result와 final_status가 "
        "정확히 일치한 비율로 계산합니다."
    )


# =========================================================
# 13. 하단 제어
# =========================================================
st.divider()

if st.button(
    "🔄 판독 화면 다시 시작",
    use_container_width=True,
):
    st.session_state.inspection_started = False
    st.session_state.inspection_finished = False
    st.session_state.elapsed_time_sec = None
    st.session_state.selected_detail_file = None
    st.session_state.current_view = "summary"
    st.rerun()